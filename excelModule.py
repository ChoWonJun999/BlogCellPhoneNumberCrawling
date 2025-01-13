import datetime
import openpyxl

import VO


class main():
    def __init__(self):
        self.wb = openpyxl.Workbook()
        # openpyxl 하면 생기는 첫번쨰 sheet 삭제
        default_sheet = self.wb.active
        self.wb.remove(default_sheet)
        # 삭제후 sheet새로생성
        self.ws = self.wb.create_sheet('Crawling')
        self.ws["A1"] = "5일 평균 방문자 수"
        self.ws["B1"] = "오늘 방문자 수"
        self.ws["C1"] = "연락처"
        self.ws["D1"] = "이웃 수"
        self.ws["E1"] = "관련도순/최신순"
        self.ws["F1"] = "키워드"
        self.ws["G1"] = "아이디"
        self.ws["O1"] = "블로그 주소"
        self.a_col_width = len(str(self.ws["A1"]))
        self.b_col_width = len(str(self.ws["B1"]))
        self.c_col_width = len(str(self.ws["C1"]))
        self.d_col_width = len(str(self.ws["D1"]))
        self.e_col_width = len(str(self.ws["E1"]))
        self.f_col_width = len(str(self.ws["F1"]))
        self.g_col_width = len(str(self.ws["G1"]))
        self.h_col_width = len(str(self.ws["O1"]))
        # 스타일 정의 keyword ROW 에 style 입히기
        # 스타일 정의
        # thick_side = Side(border_style="thick", color="000000")
        # self.thick_border = NamedStyle(name="thick_border", border=Border(left=thick_side, right=thick_side, top=thick_side, bottom=thick_side))

    def excelWrite(self):
        self.ws[f"A{VO.blogger_count}"] = VO.visitor_5
        self.ws[f"B{VO.blogger_count}"] = VO.visitor_today
        self.ws[f"C{VO.blogger_count}"] = VO.cleaned_phone_number
        self.ws[f"D{VO.blogger_count}"] = VO.blog_buddy
        self.ws[f"E{VO.blogger_count}"] = VO.search_option_value
        self.ws[f"F{VO.blogger_count}"] = VO.keyword
        self.ws[f"G{VO.blogger_count}"] = VO.blog_id
        self.ws[f"O{VO.blogger_count}"] = VO.blog_url
        self.a_col_width = max(self.a_col_width, len(str(VO.visitor_5)))
        self.b_col_width = max(self.b_col_width, len(str(VO.visitor_today)))
        self.c_col_width = max(self.c_col_width, len(str(VO.cleaned_phone_number)))
        self.d_col_width = max(self.d_col_width, len(str(VO.blog_buddy)))
        self.e_col_width = max(self.e_col_width, len(str(VO.search_option_value)))
        self.f_col_width = max(self.f_col_width, len(str(VO.keyword)))
        self.g_col_width = max(self.g_col_width, len(str(VO.blog_id)))
        self.h_col_width = max(self.h_col_width, len(str(VO.blog_url)))

    def excelSize(self):
        try:
            self.ws.column_dimensions["A"].width = max(self.a_col_width + 2, 10)
            self.ws.column_dimensions["B"].width = max(self.b_col_width + 2, 10)
            self.ws.column_dimensions["C"].width = max(self.c_col_width + 2, 10)
            self.ws.column_dimensions["D"].width = max(self.d_col_width + 2, 10)
            self.ws.column_dimensions["E"].width = max(self.e_col_width + 2, 10)
            self.ws.column_dimensions["F"].width = max(self.f_col_width + 2, 10)
            self.ws.column_dimensions["G"].width = max(self.g_col_width + 2, 10)
            self.ws.column_dimensions["O"].width = max(self.h_col_width + 2, 10)
        except Exception as e:
            print(e)

    def excelSave(self):
        now = datetime.datetime.today()
        formatted_time = now.strftime("%H%M%S")
        self.wb.save(f'{VO.save_path_value}/{VO.keyword}_{formatted_time}.xlsx')
